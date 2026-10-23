import hashlib
import io
import math
import re

import pytest

from app import create_app
from app.errors import AIError

PAGE1 = ("MASTER SERVICES AGREEMENT\n"
         "1. Term and Termination. Either party may terminate this Agreement for convenience upon thirty (30) days written notice to the other party.\n"
         "2. Confidentiality. Each party shall keep the Confidential Information of the other party strictly confidential for five (5) years.")
PAGE2 = ("3. Indemnity. The Supplier shall indemnify and hold harmless the Customer from all third-party claims arising from the Supplier's negligence.\n"
         "4. Limitation of Liability. The Supplier's aggregate liability shall not exceed the fees paid in the preceding twelve (12) months.\n"
         "5. Governing Law. This Agreement is governed by the laws of the State of New York.")
Q_TERMINATE = "Either party may terminate this Agreement for convenience upon thirty (30) days written notice"
Q_LIABILITY = "aggregate liability shall not exceed the fees paid in the preceding twelve (12) months"


def make_pdf(pages):
    """Tiny hand-rolled PDF writer (one Helvetica text stream per page)."""
    def esc(s):
        return s.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")

    objs = [b"<< /Type /Catalog /Pages 2 0 R >>", None, b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    kids = []
    for text in pages:
        page_id = len(objs) + 1
        kids.append(f"{page_id} 0 R")
        lines = ["BT /F1 9 Tf 40 760 Td 14 TL"] + [f"({esc(ln)}) '" for ln in text.split("\n")] + ["ET"]
        stream = "\n".join(lines).encode("latin-1", "replace")
        objs.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 3 0 R >> >> "
                    f"/Contents {page_id + 1} 0 R >>".encode())
        objs.append(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream")
    objs[1] = f"<< /Type /Pages /Kids [{' '.join(kids)}] /Count {len(kids)} >>".encode()
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for i, o in enumerate(objs, 1):
        offsets.append(out.tell())
        out.write(f"{i} 0 obj\n".encode() + o + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode())
    for off in offsets:
        out.write(f"{off:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    return out.getvalue()


def make_docx(paragraphs):
    import docx
    d = docx.Document()
    for p in paragraphs:
        d.add_paragraph(p)
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text = "Payment"
    t.rows[0].cells[1].text = "Net 30 days from invoice date"
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


class FakeAI:
    """Deterministic stand-in for Gemini used in tests."""
    chat_model = "fake-chat"
    embed_model = "fake-embed"
    dim = 64

    def __init__(self):
        self.handler = None
        self.fail_embed = False
        self.fail_generate = False
        self.prompts = []

    def embed(self, texts, kind="document"):
        if self.fail_embed:
            raise AIError("Gemini is rate-limiting requests. Wait a moment and retry.")
        out = []
        for t in texts:
            v = [0.0] * self.dim
            for w in re.findall(r"[a-z]{3,}", t.lower()):
                v[int(hashlib.md5(w.encode()).hexdigest(), 16) % self.dim] += 1.0
            n = math.sqrt(sum(x * x for x in v)) or 1.0
            out.append([x / n for x in v])
        return out

    def generate_json(self, system, prompt):
        self.prompts.append((system, prompt))
        if self.fail_generate:
            raise AIError("Could not reach Gemini. Check your internet connection and try again.")
        if self.handler:
            return self.handler(system, prompt)
        if "<sources>" in prompt:
            m = re.search(r'<source id="S1"[^>]*>\n(.*?)\n</source>', prompt, re.S)
            quote = m.group(1)[:80]
            return {"sufficient": True, "answer": "The document states something relevant [S1].",
                    "citations": [{"source": "S1", "quote": quote}]}
        chunks = re.findall(r'<chunk id="(C\d+)"[^>]*>\n(.*?)\n</chunk>', prompt, re.S)
        cid, text = chunks[0]
        cite = [{"source": cid, "quote": text[:60]}]
        if "Summarise" in prompt:
            return {"overview": "An agreement.", "parties": ["Supplier", "Customer"], "important_dates": [],
                    "key_points": [{"text": "Has a term.", "citations": cite},
                                   {"text": "Unsupported claim.", "citations": [{"source": cid, "quote": "this text is not in the document"}]}]}
        if "key clauses" in prompt:
            return {"clauses": [{"type": "termination", "found": True, "title": "Termination", "summary": "30 days", "citations": cite},
                                {"type": "indemnity", "found": True, "title": "Indemnity", "summary": "fake", "citations": [{"source": "C99999", "quote": "nonexistent source"}]}]}
        return {"risks": [{"title": "Short notice", "severity": "high", "category": "termination", "explanation": "x", "recommendation": "y", "citations": cite},
                          {"title": "Invented", "severity": "critical", "category": "z", "explanation": "x", "recommendation": "y", "citations": []}]}


@pytest.fixture()
def fake():
    return FakeAI()


def _make(tmp_path, **kw):
    cfg = {"DATA_DIR": str(tmp_path), "PROCESS_SYNC": True, "RATELIMIT_ENABLED": False, "GEMINI_API_KEY": "",
           "SECRET_KEY": "test-secret"}
    cfg.update(kw)
    return create_app(cfg)


@pytest.fixture()
def app(tmp_path, fake):
    a = _make(tmp_path)
    a.extensions["lexai_ai"] = fake
    return a


@pytest.fixture()
def app_no_key(tmp_path):
    return _make(tmp_path)


class Api:
    def __init__(self, app):
        self.app = app
        self.c = app.test_client()
        self.token = self.c.get("/api/auth/csrf").get_json()["csrf_token"]

    def req(self, method, url, **kw):
        headers = kw.pop("headers", {})
        headers.setdefault("X-CSRF-Token", self.token)
        r = self.c.open(url, method=method, headers=headers, **kw)
        try:
            j = r.get_json(silent=True)
        except Exception:
            j = None
        if isinstance(j, dict) and j.get("csrf_token"):
            self.token = j["csrf_token"]
        r.j = j
        return r

    def get(self, url, **kw):
        return self.req("GET", url, **kw)

    def post(self, url, json=None, **kw):
        return self.req("POST", url, json=json, **kw)

    def patch(self, url, json=None, **kw):
        return self.req("PATCH", url, json=json, **kw)

    def delete(self, url, **kw):
        return self.req("DELETE", url, **kw)

    def register(self, email="a@example.com", password="Passw0rd-long", name="Alice"):
        r = self.post("/api/auth/register", {"email": email, "password": password, "name": name})
        assert r.status_code == 201, r.j
        return r

    def upload(self, data, filename="contract.pdf", url="/api/documents", **form):
        return self.req("POST", url, data={"file": (io.BytesIO(data), filename), **form},
                        content_type="multipart/form-data")


@pytest.fixture()
def api(app):
    a = Api(app)
    a.register()
    return a


@pytest.fixture()
def api_factory():
    return Api


@pytest.fixture()
def pdf_bytes():
    return make_pdf([PAGE1, PAGE2])


@pytest.fixture()
def doc(api, pdf_bytes):
    r = api.upload(pdf_bytes, title="Services Agreement")
    assert r.status_code == 202, r.j
    return r.j["document"]
